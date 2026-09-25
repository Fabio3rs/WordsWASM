"""Generate a local, reproducible before/after rejection workbook."""

from __future__ import annotations

import argparse
from collections import Counter
import json
from pathlib import Path
import subprocess
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


TEXT = (
    "Quae res in Civitate duae plurimum possunt, eae contra nos ambae "
    "faciunt in hoc tempore, summa gratia et eloquentia;"
)


def run(binary: Path, *arguments: str) -> dict:
    completed = subprocess.run(
        [str(binary), "--text", TEXT, "--strategy", "dependency-mst",
         "--json", *arguments],
        check=True,
        capture_output=True,
        text=True,
    )
    return json.loads(completed.stdout)


def style_sheet(sheet, widths: list[int]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--binary", type=Path,
        default=Path("build/parsers/parsers_investigation/parsers_investigation"),
    )
    parser.add_argument(
        "--output", type=Path,
        default=Path("parsers_investigation/reports/quae_res_rejeicoes.xlsx"),
    )
    args = parser.parse_args()
    binary = args.binary.resolve()

    # An explicit corpus fixture takes the exact, pre-fallback path while
    # preserving the text, grammar mode, strategy, data, and product budget.
    with tempfile.TemporaryDirectory() as directory:
        corpus = Path(directory) / "baseline.json"
        corpus.write_text(
            json.dumps({
                "schema": "words-parser-fixtures",
                "schemaVersion": 2,
                "fixtures": [{
                    "id": "spreadsheet-baseline",
                    "text": TEXT,
                    "phenomenon": "local-comparison",
                    "mode": "complete-clause",
                }],
            }),
            encoding="utf-8",
        )
        before = run(binary, "--corpus", str(corpus))
    after = run(binary, "--include-rejections")

    if before["status"] != "experiment-budget-exceeded":
        raise RuntimeError("baseline did not stop at the product budget")
    if after["status"] != "approximate":
        raise RuntimeError("local beam search was not used")

    rejected = after["enumeration"]["rejectedAlternatives"]
    grammatical = after["enumeration"]["rejectionsByConstraint"]
    score_cut = after["enumeration"]["beamScorePrunedStates"]
    if len(rejected) != sum(grammatical.values()) + score_cut:
        raise RuntimeError("rejection detail does not match aggregate counts")

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumo"
    summary.append(["Métrica", "Antes: limite exato", "Depois: feixe", "Interpretação"])
    metrics = [
        ("Estado", before["status"], after["status"], "Depois é aproximado"),
        ("Produto morfológico bruto", int(before["morphology"]["rawProduct"]),
         int(after["morphology"]["rawProduct"]), "Combinações possíveis, não estados visitados"),
        ("Produto após propagação", int(before["propagation"]["prunedProduct"]),
         int(after["propagation"]["prunedProduct"]), "Nenhuma redução nesta frase"),
        ("Estados parciais visitados", before["enumeration"]["partialStates"],
         after["enumeration"]["partialStates"], "Inclui tentativas aceitas e rejeitadas"),
        ("Rejeitados por gramática", sum(before["enumeration"]["rejectionsByConstraint"].values()),
         sum(grammatical.values()), "H005: complemento possível para preposição"),
        ("Descartados pelo score", before["enumeration"].get("beamScorePrunedStates", 0),
         score_cut, "Podem ser gramaticalmente válidos"),
        ("Combinações completas examinadas", before["enumeration"]["completeAssignments"],
         after["enumeration"]["completeAssignments"], "Após a poda do feixe"),
        ("Análises morfológicas aceitas", before["acceptance"]["morphAssignments"],
         after["acceptance"]["morphAssignments"], "Antes não chegou à enumeração"),
        ("Árvores decodificadas", before["decoder"]["completeTrees"],
         after["decoder"]["completeTrees"], "Não usadas para a lista de rejeições"),
    ]
    for row in metrics:
        summary.append(row)
    style_sheet(summary, [38, 25, 25, 65])

    reasons = workbook.create_sheet("Motivos")
    reasons.append(["Motivo", "Antes", "Depois", "Tipo"])
    for reason, count in sorted(grammatical.items()):
        reasons.append([reason, before["enumeration"]["rejectionsByConstraint"].get(reason, 0),
                        count, "Regra gramatical"])
    reasons.append(["beam-score-limit", 0, score_cut, "Corte por score; não prova invalidade"])
    style_sheet(reasons, [30, 15, 15, 45])

    grouped = Counter()
    example_prefix = {}
    for item in rejected:
        choice = item["alternative"]
        key = (
            item["reason"], choice["token"], choice["candidate"],
            choice["lemma"], choice["part"], choice["morphology"],
        )
        grouped[key] += 1
        example_prefix.setdefault(key, ",".join(
            str(index + 1) for index in item["prefixCandidates"]
        ))

    alternatives = workbook.create_sheet("Alternativas")
    alternatives.append([
        "Motivo", "Token", "Forma", "Candidato", "Lema", "Classe",
        "Morfologia", "Ocorrências", "Exemplo de prefixo (candidatos 1-based)",
    ])
    surfaces = after["morphology"]["surfaceTokens"]
    for key, count in sorted(grouped.items()):
        reason, token, candidate, lemma, part, morphology = key
        alternatives.append([
            reason, token + 1, surfaces[token], candidate + 1, lemma,
            part, morphology, count, example_prefix[key],
        ])
    style_sheet(alternatives, [23, 10, 18, 12, 23, 18, 45, 14, 75])

    states = workbook.create_sheet("Estados rejeitados")
    states.append([
        "Nº", "Motivo", "Token", "Forma", "Candidato", "Lema", "Classe",
        "Morfologia", "Score parcial", "Prefixo (candidatos 1-based)",
    ])
    for number, item in enumerate(rejected, start=1):
        choice = item["alternative"]
        states.append([
            number, item["reason"], choice["token"] + 1,
            surfaces[choice["token"]], choice["candidate"] + 1,
            choice["lemma"], choice["part"], choice["morphology"],
            item["partialScore"], ",".join(
                str(index + 1) for index in item["prefixCandidates"]
            ),
        ])
    style_sheet(states, [10, 23, 10, 18, 12, 23, 18, 45, 16, 95])

    notes = workbook.create_sheet("Notas")
    notes.append(["Campo", "Descrição"])
    notes.append(["Frase", TEXT])
    notes.append(["Estratégia", "dependency-mst; as rejeições listadas precedem a construção de árvores"])
    notes.append(["Antes", "Mesma estratégia, texto e base, com fixture explícita para reproduzir o caminho exato sem feixe"])
    notes.append(["Limite", before["maxProduct"]])
    notes.append(["Base", after["datasetId"]])
    notes.append(["H005", "Preposição sem complemento compatível possível no prefixo examinado"])
    notes.append(["beam-score-limit", "Prefixo saiu dos 512 melhores por score; pode ser gramaticalmente válido"])
    notes.append(["Ocorrências", "A mesma alternativa pode ser rejeitada em um contexto e sobreviver em outro"])
    notes.append(["Prefixo", "Índices dos candidatos escolhidos, da primeira palavra até o token da linha"])
    notes.append(["Limitação", "A lista contém apenas estados visitados pelo feixe, não todas as combinações possíveis"])
    style_sheet(notes, [25, 110])

    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    print(f"{args.output}: {len(rejected)} estados, {len(grouped)} alternativas agrupadas")


if __name__ == "__main__":
    main()
